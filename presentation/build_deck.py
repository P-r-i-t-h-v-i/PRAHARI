"""Build PRAHARI's six-slide SIH deck from the supplied SIH template."""
from pathlib import Path
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.oxml.xmlchemy import OxmlElement
from PIL import Image
import json

ROOT=Path(__file__).resolve().parent
A=ROOT/'assets'
TEMPLATE=Path('/home/per/Downloads/SIH2026-IDEA-Presentation-Format.pptx')
prs=Presentation(TEMPLATE)
N='152C3B'; IV='F6F4EE'; AM='E5AE48'; TE='287F79'; MUT='5B6C73'; LINE='D9DFDC'; WH='FFFFFF'; PALE='E7EFEB'
FONT='Lato'

def rect(s,x,y,w,h,c,r=False,line=None):
    sh=s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE if r else MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    sh.fill.solid();sh.fill.fore_color.rgb=RGBColor.from_string(c)
    if line: sh.line.color.rgb=RGBColor.from_string(line);sh.line.width=Pt(.7)
    else: sh.line.fill.background()
    if r: sh.adjustments[0]=.08
    sh._element.spPr.append(OxmlElement('a:effectLst'))
    for ref in sh._element.xpath('.//a:effectRef'): ref.set('idx','0')
    return sh

def txt(s,text,x,y,w,h,size=18,color=N,bold=False,font=FONT,align=None,url=None):
    sh=s.shapes.add_textbox(Inches(x),Inches(y),Inches(w),Inches(h))
    tf=sh.text_frame;tf.clear();tf.word_wrap=True
    tf.margin_left=tf.margin_right=0;tf.margin_top=tf.margin_bottom=0
    for i,line in enumerate(text.split('\n')):
        p=tf.paragraphs[0] if i==0 else tf.add_paragraph()
        p.text=line;p.font.name=font;p.font.size=Pt(size);p.font.bold=bold;p.font.color.rgb=RGBColor.from_string(color)
        p.space_before=Pt(0);p.space_after=Pt(2);p.line_spacing=1.05
        if align is not None:p.alignment=align
        if url:
            sh.click_action.hyperlink.address=url
            for run in p.runs:
                run.font.color.rgb=RGBColor.from_string(color)
                run.font.underline=False
    return sh

def rule(s,x,y,w,c=LINE):rect(s,x,y,w,.014,c)

def image_fit(s,path,x,y,w,h):
    im=Image.open(path);ratio=im.width/im.height
    if ratio>w/h:hh=w/ratio;ww=w
    else:hh=h;ww=h*ratio
    return s.shapes.add_picture(str(path), Inches(x+(w-ww)/2), Inches(y+(h-hh)/2), width=Inches(ww),height=Inches(hh))

def photo(s,path,x,y,w,h):
    im=Image.open(path);ratio=im.width/im.height
    p=s.shapes.add_picture(str(path),Inches(x),Inches(y),width=Inches(w),height=Inches(h))
    if ratio>w/h:
        crop=(1-(w/h)/ratio)/2;p.crop_left=crop;p.crop_right=crop
    else:
        crop=(1-ratio/(w/h))/2;p.crop_top=crop;p.crop_bottom=crop
    return p

def tag(s,text,x,y,w,c=TE):
    rect(s,x,y,w,.3,c,True);txt(s,text,x+.08,y+.05,w-.16,.22,9,WH,True)

def line(s,x1,y1,x2,y2,c=TE,width=2):
    sh=s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT,Inches(x1),Inches(y1),Inches(x2),Inches(y2));sh.line.color.rgb=RGBColor.from_string(c);sh.line.width=Pt(width)

def arrow(s,x,y,w=.35,c=TE):
    sh=s.shapes.add_shape(MSO_SHAPE.CHEVRON,Inches(x),Inches(y),Inches(w),Inches(.23));sh.fill.solid();sh.fill.fore_color.rgb=RGBColor.from_string(c);sh.line.fill.background()

def base(s,i,title):
    # Retain the supplied presentation package, slide order, dimensions, layouts,
    # and masters. Replace placeholder contents with designed SIH sections.
    for sh in list(s.shapes):sh._element.getparent().remove(sh._element)
    s.background.fill.solid();s.background.fill.fore_color.rgb=RGBColor.from_string(IV)
    rect(s,0,0,13.333,1.14,WH)
    rect(s,.4,.32,1.46,.47,N,True)
    txt(s,'POINT BREAK',.48,.465,1.30,.23,10,WH,True,align=PP_ALIGN.CENTER)
    txt(s,title,2.05,.38,8.6,.50,23,N,True,align=PP_ALIGN.CENTER)
    image_fit(s,A/'template_2.png',10.95,.17,1.94,.80)
    rule(s,.4,1.12,12.5)
    rect(s,0,6.95,13.333,.55,N)
    txt(s,'PRAHARI  /  POINT BREAK',.45,7.12,3,.2,9,WH,True)
    txt(s,'@SIH Idea submission- Template',4.8,7.12,4.1,.2,9,'CFD9DB',align=PP_ALIGN.CENTER)
    txt(s,f'{i:02d} / 06',11.65,7.10,1.15,.23,10,WH,True,align=PP_ALIGN.RIGHT)

def note(s,text):s.notes_slide.notes_text_frame.text=text

def bullet(s,title,body,x,y,w=3.6,color=TE,size=15):
    rect(s,x,y+.065,.07,.23,color)
    txt(s,title,x+.18,y,w-.18,.34,size,N,True)
    txt(s,body,x+.18,y+.4,w-.18,.65,size-1,MUT)

# Delete the instruction-only slide, as explicitly directed by the template.
sid=prs.slides._sldIdLst[-1];prs.part.drop_rel(sid.rId);prs.slides._sldIdLst.remove(sid)

# 01 — title + mandatory registration fields
s=prs.slides[0];base(s,1,'SMART INDIA HACKATHON 2026')
rect(s,.45,1.38,7.18,2.38,N,True)
txt(s,'PRAHARI',.72,1.52,6.6,.94,57,WH,True)
txt(s,'THE COUNTDOWN BAND',.76,2.53,6.25,.4,19,AM,True)
txt(s,'The hazard outside. The strain inside.\nOne warning on your wrist.',.76,3.02,6.35,.61,19,WH)
photo(s,A/'workers_mumbai.jpg',7.86,1.38,5.02,4.93)
rect(s,7.86,5.35,5.02,.96,N)
txt(s,'Built for people\nworking beyond the last signal.',8.1,5.49,4.55,.68,20,WH,True)
txt(s,'Photo: yajna S R / Pexels · Mumbai, India',7.9,6.41,4.95,.22,9,MUT)
tag(s,'PROBLEM STATEMENT ID  26181',.49,3.98,3.17,N)
txt(s,'PROBLEM STATEMENT TITLE',.5,4.43,6.8,.22,10,TE,True)
ps='A secure, AI-powered Personal Health Companion that delivers real-time, privacy-preserving health monitoring and early warning capabilities, helping individuals recognize health risks before they become emergencies. The solution should improve resilience during heat waves, floods, pollution events, and other disasters common in India while enabling continuous health support through on-device intelligence.'
txt(s,ps,.5,4.77,7.0,1.39,13.2,N)
rule(s,.5,6.14,7.0)
txt(s,'THEME',.5,6.28,3.5,.19,9,TE,True)
txt(s,'MedTech / BioTech / HealthTech',.5,6.49,3.8,.23,12,N,True)
txt(s,'PS CATEGORY',4.48,6.28,1.25,.19,9,TE,True)
txt(s,'Hardware',4.48,6.49,1.2,.23,12,N,True)
txt(s,'TEAM ID',6.0,6.28,1.5,.19,9,TE,True)
txt(s,'To be provided',6.0,6.49,1.55,.23,11,N,True)
note(s,'Opening: PRAHARI aims to bring two early warnings to one wrist: official hazard alerts and personal heat-strain trends. Team name and PS ID/title are taken from PRAHARI_SIH2026_Idea_3.pdf. Team ID was absent from the supplied materials and must be filled before submission. The photograph is a real contextual photograph, not the team or a prototype trial. Photo: yajna S R, https://www.pexels.com/photo/construction-workers-working-on-top-of-a-building-20591230/ (Pexels License).')

# 02 — proposed solution, problem and innovation
s=prs.slides[1];base(s,2,'IDEA TITLE')
txt(s,'Two threats. One personal warning system.',.49,1.37,12.2,.5,30,N,True)
txt(s,'Proposed Solution  /  PRAHARI — The Countdown Band',.5,1.95,12,.3,14,MUT)
for x,num,title,sub,col in [(.5,'01','THE HAZARD OUTSIDE','Warnings must travel past a network outage.',TE),(6.85,'02','THE STRAIN INSIDE','A regional forecast cannot see your body.',AM)]:
    rect(s,x,2.48,5.98,1.84,WH,True)
    rect(s,x,2.48,.065,1.84,col)
    txt(s,num,x+.22,2.66,.63,.6,30,col,True)
    txt(s,title,x+1,2.67,4.7,.34,19,N,True)
    txt(s,sub,x+.23,3.18,5.47,.4,16,MUT)
txt(s,'Official CAP alerts → phone + gateway band\n→ LoRa peer relay → vibration at the wrist',.74,3.70,5.45,.51,14,N,True)
txt(s,'Heart rate + skin temperature + ambient heat\n→ on-device trend → early rest / cooling prompt',7.08,3.70,5.48,.51,14,N,True)
rect(s,.5,4.60,12.33,1.14,N,True)
txt(s,'THE BAND',.74,4.80,2.0,.28,14,AM,True)
txt(s,'Senses • decides • warns',.74,5.17,3.35,.30,16,WH)
line(s,4.2,4.82,4.2,5.48,'506571',1)
txt(s,'THE COMPANION APP',4.46,4.80,3.5,.28,14,AM,True)
txt(s,'Fetches alerts • keeps offline history',4.46,5.17,4.52,.30,16,WH)
txt(s,'ONE CONNECTED ENTRY POINT\nSHARES WARNINGS LOCALLY',9.43,4.97,3.1,.5,12,WH,True)
txt(s,'INNOVATION & UNIQUENESS',.52,6.02,3.2,.28,12,TE,True)
txt(s,'Body + hazard in one device',.52,6.42,3.9,.28,17,N,True)
txt(s,'Time to heat-risk threshold*',4.76,6.42,4.1,.28,17,N,True)
txt(s,'Recovery across shifts*',9.02,6.42,3.85,.28,17,N,True)
txt(s,'*Proposed features; personal forecasts and recovery tracking require validation.',4.76,6.06,8.1,.25,10,MUT)
note(s,'Detailed solution: A connected phone downloads official CAP alerts and sends relevant alerts over BLE to a gateway band. That LoRa-equipped band can relay them to nearby bands. Phones alone have no LoRa radio. New external alerts need at least one connected entry point; local body monitoring is intended to continue offline. Proposed originality is the combination of official multi-hazard warning, local relaying, individualized heat-strain forecasting, and a thermal-debt/recovery ledger. Forecast-to-threshold is a research feature, not a validated countdown to collapse. Do not present this as implemented mesh or proven prediction. SACHET confirms CAP-based multi-hazard dissemination and RSS: https://sachet.ndma.gov.in/.')

# 03 — technical architecture
s=prs.slides[2];base(s,3,'TECHNICAL APPROACH')
txt(s,'Sense locally. Decide on-device. Relay the warning.',.5,1.36,12.3,.51,28,N,True)
txt(s,'Proposed rev B architecture · hardware + embedded intelligence + offline companion',.5,1.96,12,.30,14,MUT)
for x,num,title,col in [(.5,'1','SENSE',TE),(4.71,'2','DECIDE',TE),(8.92,'3','WARN',AM)]:
    rect(s,x,2.48,3.9,2.34,WH,True)
    rect(s,x+.18,2.68,.47,.47,col,True)
    txt(s,num,x+.18,2.76,.47,.26,17,WH,True,align=PP_ALIGN.CENTER)
    txt(s,title,x+.81,2.77,2.9,.4,20,N,True)
txt(s,'MAX30102 · heart rate / PPG\nMAX30208 · contact skin temp\nAHT20 · air temp + humidity\nOnboard IMU · motion context',.73,3.39,3.46,1.12,15,N)
txt(s,'Signal-quality + motion gate\nKalman estimate + personal baseline\nTrend-to-threshold forecast*\nRecovery / thermal-debt ledger*',4.94,3.39,3.45,1.16,14.2,N)
txt(s,'Distinct haptics + OLED status\nSX1262 LoRa · local peer relay*\nBLE · Android companion*\nLocal records + protected keys*',9.15,3.39,3.45,1.16,14.5,N)
arrow(s,4.44,3.62,.22);arrow(s,8.65,3.62,.22)
image_fit(s,A/'xiao_sense.png',.55,4.93,1.33,1.06)
txt(s,'XIAO nRF52840 Sense',2.01,4.98,4.55,.3,17,N,True)
txt(s,'C/C++ · Arduino   |   TinyML planned\nFlutter / Dart · local storage planned',2.01,5.39,5.33,.60,14,MUT)
txt(s,'Privacy by design',8.0,4.98,4.43,.3,17,N,True)
txt(s,'Keep raw physiology on-device;\nshare only necessary alert data.*',8.0,5.39,4.68,.60,14,MUT)
rect(s,.5,6.17,12.33,.48,PALE,True)
txt(s,'CAP → connected phone → BLE gateway band → LoRa peers → haptic warning',.7,6.30,11.91,.27,16,TE,True,align=PP_ALIGN.CENTER)
txt(s,'Board image: Seeed Studio   ·   *Planned / requires validation; air temperature + humidity alone are not full WBGT.',.56,6.72,12.2,.18,9,MUT)
note(s,'Methodology: sample sensors on the shared I2C bus; use motion and signal quality to decide whether the physiological estimate is usable; build individual baselines; evaluate trend prediction before presenting a time estimate. Low-quality input must show an unavailable estimate, not reassuring status. Rev B uses MAX30102, MAX30208 and SX1262/Ra-01SH from pcb/README.md and the project report; older breadboard firmware uses SX1278 and an earlier sensor breakout. LoRa relay, TinyML, app, security integration and forecasting are proposed, not demonstrated by the checked-in sensor test firmware. AHT20 measures air temperature and humidity, not solar radiant load; do not call this full measured WBGT. Sources: https://wiki.seeedstudio.com/XIAO_BLE/ ; https://ojs.aaai.org/index.php/AAAI/article/view/18809 ; https://sachet.ndma.gov.in/ . Manufacturer image shows an actual XIAO module, not a completed PRAHARI band.')

# 04 — feasibility, risk and roadmap
s=prs.slides[3];base(s,4,'FEASIBILITY AND VIABILITY')
txt(s,'A working sensing core. A focused path to a field pilot.',.5,1.36,12.3,.52,28,N,True)
rect(s,.5,2.18,4.0,3.39,N,True)
txt(s,'FEASIBILITY TODAY',.75,2.42,3.46,.34,17,AM,True)
txt(s,'Verified on the bench',.75,2.97,3.5,.32,18,WH,True)
txt(s,'Live heart-rate sensing + OLED\nI²C discovery and sensor diagnostics',.75,3.42,3.47,.68,15,WH)
txt(s,'Engineering in progress',.75,4.24,3.5,.34,18,WH,True)
txt(s,'Rev B PCB + BOM available\nRadio bring-up / PCB checks open\nApp + forecasting + relay to integrate',.75,4.68,3.47,.75,14,'D2DFE0')
txt(s,'CHALLENGE',4.84,2.25,2.61,.26,11,TE,True)
txt(s,'STRATEGY TO OVERCOME IT',7.68,2.25,4.99,.26,11,TE,True)
rows=[('Motion & sensor uncertainty','Quality gate; reference-device validation;\nshow “estimate unavailable” on poor signal.'),('Radio loss & stale alerts','Time stamps, expiry, deduplication, hop limits;\nmeasure delivery and battery life in the field.'),('Fit, power & deployment cost','Modular parts; duty cycling; enclosure trials;\nquote the complete rev B assembly before pilot.')]
for j,(a,b) in enumerate(rows):
    y=2.78+j*.92;rule(s,4.82,y-.11,7.97)
    txt(s,a,4.84,y,2.59,.66,16,N,True)
    txt(s,b,7.68,y,5.08,.7,15,MUT)
txt(s,'Pilot readiness gates',.52,5.83,3.5,.3,15,TE,True)
for x,n,t in [(3.92,'01','Close hardware issues'),(6.94,'02','Validate sensing + relay'),(10.05,'03','Supervised worker pilot')]:
    txt(s,n,x,5.82,.46,.36,18,TE,True)
    txt(s,t,x+.55,5.85,2.23,.55,14,N,True)
txt(s,'Economic viability: off-the-shelf modules + a two-layer PCB; final cost awaits rev B BOM, battery, enclosure and assembly quotes.',.53,6.51,12.15,.29,12,MUT)
note(s,'Evidence: README.md and Obsidian Projects/PRAHARI.md document live finger heart-rate testing and successful OLED integration. These are bench demonstrations, not clinical or field validation. The latest dated root notes record unresolved SX1278 SPI bring-up and PCB DRC issues (3 errors + 4 unconnected); the older pcb README contains a cleaner historical result. Do not claim the current board is fabrication-ready. A 70 x 40 mm rev B carrier design and BOM exist. The older ₹3,468 number covers selected rev A modules only and is not a finished rev B device price; it was omitted as a headline. The old ~2 km/hop and clinical accuracy values are unverified for this implementation. The next gates are hardware integrity, safe sensor/reference testing, radio delivery and power measurements, then supervised wearability and usability work. Applicable radio settings and permissions must be checked for the chosen region before field operation; no categorical regulatory approval is claimed.')

# 05 — impact and benefits
s=prs.slides[4];base(s,5,'IMPACT AND BENEFITS')
txt(s,'Earlier awareness. A clearer next action.',.5,1.36,12.3,.51,30,N,True)
photo(s,A/'workers_mumbai.jpg',.5,2.15,4.24,3.86)
rect(s,.5,5.02,4.24,.99,N)
txt(s,'Start with outdoor work crews.',.71,5.20,3.78,.39,19,WH,True)
txt(s,'Construction · agriculture · municipal work',.71,5.68,3.79,.24,11,WH)
txt(s,'57%',5.11,2.15,2.25,.91,52,TE,True)
txt(s,'of Indian districts face high\nto very high heat risk.',7.6,2.34,4.94,.69,21,N,True)
txt(s,'CEEW, 2025 · district-level heat-risk assessment [1]',5.17,3.24,7.53,.25,11,MUT)
rule(s,5.15,3.73,7.63)
bullet(s,'SOCIAL  /  Reach the person','Simple haptics; no personal smartphone needed\nfor local warning reception.',5.15,3.98,7.61,TE,16)
bullet(s,'ECONOMIC  /  Support timely action','Earlier rest and cooling prompts aim to reduce\ndisruption; pilot outcomes will test the benefit.',5.15,5.15,7.61,TE,16)
rect(s,.5,6.29,12.33,.44,PALE,True)
txt(s,'MEASURE IN THE PILOT   Alert delivery time  •  Missed / false alerts  •  Time to action  •  Shift-long wearability',.72,6.42,11.94,.24,12,TE,True)
txt(s,'Photo: yajna S R / Pexels',.53,6.07,4.17,.18,9,MUT)
note(s,'Potential impact, not demonstrated outcomes. Intended initial audience is outdoor work crews, followed by field responders and remote communities. Simple haptics and icon/status displays can reduce dependence on text literacy and individual phone ownership; validate distinct patterns and comprehension in the user pilot. Crew response still requires training and a local escalation protocol. Economic savings and reduced health harm are not quantified until evaluated. Modular repair and rechargeable power are design intentions; no environmental benefit is quantified. CEEW reports 57% of Indian districts at high to very high heat risk in its 2025 study: https://www.ceew.in/publications/mapping-climate-risks-and-impacts-of-extreme-heatwave-disaster-in-indian-districts . Photo is contextual, not an endorsement or completed trial: https://www.pexels.com/photo/construction-workers-working-on-top-of-a-building-20591230/ .')

# 06 — research references and project evidence
s=prs.slides[5];base(s,6,'RESEARCH AND REFERENCES')
txt(s,'Grounded in research. Built to be tested.',.5,1.36,12.2,.51,30,N,True)
txt(s,'Reference & research work · clickable source titles',.52,1.99,12.1,.3,14,MUT)
refs=[
('01','THE NEED','CEEW · How Extreme Heat is Impacting India (2025)','District-level heat risk; basis for the 57% figure.', 'https://www.ceew.in/publications/mapping-climate-risks-and-impacts-of-extreme-heatwave-disaster-in-indian-districts'),
('02','OFFICIAL WARNING SOURCE','NDMA · SACHET National Disaster Alert Portal','CAP-based multi-hazard alerts and RSS dissemination.', 'https://sachet.ndma.gov.in/'),
('03','ALGORITHM FOUNDATION','Buller et al. · Wearable physiological estimation','Kalman-filter research; PRAHARI still needs its own validation.', 'https://ojs.aaai.org/index.php/AAAI/article/view/18809'),
('04','HARDWARE PLATFORM','Seeed Studio · XIAO nRF52840 Series','MCU, BLE and Sense IMU documentation; module image source.', 'https://wiki.seeedstudio.com/XIAO_BLE/'),
]
for j,(n,k,t,d,url) in enumerate(refs):
    y=2.61+j*.96
    txt(s,n,.53,y,.52,.44,24,TE,True)
    txt(s,k,1.25,y,7.1,.23,10,TE,True)
    txt(s,t,1.25,y+.28,7.12,.31,15.5,N,True,url=url)
    txt(s,d,1.25,y+.65,7.1,.26,12,MUT)
    if j<3:rule(s,.55,y+.89,7.96)
rect(s,8.91,2.58,3.92,3.65,N,True)
txt(s,'PROJECT EVIDENCE',9.16,2.82,3.42,.29,15,AM,True)
txt(s,'PRAHARI repository',9.16,3.34,3.42,.33,18,WH,True,url='https://github.com/P-r-i-t-h-v-i/PRAHARI')
txt(s,'Sensor + OLED firmware\nKiCad design + BOM\nBring-up and verification notes',9.16,3.84,3.40,.98,15,WH)
txt(s,'Obsidian project brain\nPRAHARI reference deck + report',9.16,5.03,3.40,.7,14,'D2DFE0')
txt(s,'Repository access may be restricted.',9.16,5.83,3.36,.22,10,'D2DFE0')
txt(s,'PHOTOGRAPHY  yajna S R / Pexels · Mumbai construction workers · Pexels License',.54,6.57,12.1,.24,11,MUT,url='https://www.pexels.com/photo/construction-workers-working-on-top-of-a-building-20591230/')
note(s,'Local reference materials: /home/per/Downloads/SIH2026-IDEA-Presentation-Format.pptx; /home/per/Downloads/PRAHARI_SIH2026_Idea_3.pdf; /home/per/Downloads/PRAHARI_Project_Report.pdf; /home/per/Documents/Obsidian/Projects/PRAHARI.md and linked XIAO, LoRa and MAX30102 notes; /home/per/PRAHARI/README.md; /home/per/PRAHARI/pcb/README.md. The supplied template is the base package; all six idea-detail sections are retained and the instruction slide is removed. Research supports motivation and architecture, not achieved PRAHARI performance. No mortality guarantee, regulatory approval, clinical accuracy, competitor claim, field range, or manufactured device cost is asserted. Real stock photography only; the board image is from its manufacturer. No AI image generation used. Photo license: https://www.pexels.com/license/.')

prs.core_properties.title='PRAHARI — The Countdown Band | SIH 2026'
prs.core_properties.subject='PS 26181 · Point Break · Six-slide SIH idea presentation'
prs.core_properties.author='Point Break'
prs.core_properties.keywords='PRAHARI, SIH 2026, Point Break, early warning, wearable, LoRa'
prs.save(ROOT/'PRAHARI_SIH2026.pptx')
issues=[]
for i,s in enumerate(prs.slides,1):
    for sh in s.shapes:
        if sh.left<0 or sh.top<0 or sh.left+sh.width>prs.slide_width+1000 or sh.top+sh.height>prs.slide_height+1000:
            issues.append([i,sh.name,'outside slide'])
(ROOT/'layout_check.json').write_text(json.dumps({'slides':len(prs.slides),'out_of_bounds':issues},indent=2))
assert not issues,issues
print('Saved',ROOT/'PRAHARI_SIH2026.pptx')
